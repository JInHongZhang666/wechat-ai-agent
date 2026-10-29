from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "07-AI信息小屋Agent实训报告.docx"
ARCHITECTURE = ROOT / "01-Agent架构设计图.png"

NAVY = "17365D"
BLUE = "DCE6F1"
PALE = "F4F7FB"
GRID = "D9D9D9"
TEXT = RGBColor(0, 0, 0)


def font(size: int, bold: bool = False, color=(0, 0, 0)):
    candidates = [
        Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def draw_arrow(draw: ImageDraw.ImageDraw, start, end, color="#315B7D", width=6):
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        sign = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - sign * 22, y2 - 13), (x2 - sign * 22, y2 + 13)]
    else:
        sign = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 13, y2 - sign * 22), (x2 + 13, y2 - sign * 22)]
    draw.polygon(points, fill=color)


def centered(draw, box, title, subtitle="", fill="#FFFFFF", outline="#315B7D"):
    draw.rounded_rectangle(box, radius=24, fill=fill, outline=outline, width=4)
    x1, y1, x2, y2 = box
    title_font = font(34, True)
    sub_font = font(23)
    tb = draw.textbbox((0, 0), title, font=title_font)
    tx = (x1 + x2 - (tb[2] - tb[0])) / 2
    ty = y1 + 24
    draw.text((tx, ty), title, fill="#102A43", font=title_font)
    if subtitle:
        sb = draw.multiline_textbbox((0, 0), subtitle, font=sub_font, spacing=7, align="center")
        sx = (x1 + x2 - (sb[2] - sb[0])) / 2
        draw.multiline_text((sx, ty + 54), subtitle, fill="#486581", font=sub_font, spacing=7, align="center")


def build_architecture():
    image = Image.new("RGB", (1900, 1250), "#F7FAFC")
    draw = ImageDraw.Draw(image)
    draw.text((80, 45), "AI 信息小屋 Agent 运行架构", fill="#102A43", font=font(50, True))
    draw.text((82, 110), "LangGraph 有界 ReAct 循环与工具权限边界", fill="#486581", font=font(27))

    boxes = {
        "user": (80, 240, 350, 420),
        "thought": (480, 210, 810, 450),
        "action": (940, 210, 1270, 450),
        "observation": (1400, 210, 1810, 450),
        "answer": (740, 560, 1120, 780),
    }
    centered(draw, boxes["user"], "用户输入", "主题、时间范围\n后续修正条件", "#E6F6FF")
    centered(draw, boxes["thought"], "Thought", "结合历史与观察\n选择工具或结束", "#FFFBEA")
    centered(draw, boxes["action"], "Action", "原生 function call\nJSON Schema 校验", "#EAF7EF")
    centered(draw, boxes["observation"], "Observation", "结构化工具结果\n写回 messages", "#F3E8FF")
    centered(draw, boxes["answer"], "Answer", "基于查询结果回答\n附来源与证据限制", "#E8F1FF")

    draw_arrow(draw, (350, 330), (480, 330))
    draw_arrow(draw, (810, 330), (940, 330))
    draw_arrow(draw, (1270, 330), (1400, 330))
    draw_arrow(draw, (1605, 450), (1120, 650))
    draw_arrow(draw, (1400, 395), (810, 395))
    draw_arrow(draw, (930, 560), (930, 450))

    tool_x = [90, 440, 790, 1140, 1490]
    tools = [
        ("本地资讯搜索", "search_ai_news"),
        ("补充独立来源", "find_supporting_sources"),
        ("构建研究证据", "build_research_evidence"),
        ("授权刷新信源", "refresh_news_sources"),
        ("AIHOT 发现", "search_aihot_news"),
    ]
    for x, (name, code) in zip(tool_x, tools):
        centered(draw, (x, 920, x + 300, 1090), name, code, "#FFFFFF", "#829AB1")

    draw.text((90, 850), "工具层", fill="#102A43", font=font(34, True))
    draw.line((230, 780, 230, 920), fill="#315B7D", width=5)
    draw.line((230, 850, 1640, 850), fill="#315B7D", width=5)
    for x in tool_x:
        draw_arrow(draw, (x + 150, 850), (x + 150, 920), width=4)

    draw.rounded_rectangle((1250, 560, 1810, 780), radius=24, fill="#FFFFFF", outline="#829AB1", width=4)
    draw.text((1290, 590), "会话记忆与运行状态", fill="#102A43", font=font(32, True))
    draw.multiline_text(
        (1290, 650),
        "messages：用户与助手历史\nobservations：本轮工具观察\nseen：本轮允许进入证据链的文章 ID\n限制：进程内保存，重启后清空",
        fill="#486581",
        font=font(23),
        spacing=8,
    )
    draw_arrow(draw, (1250, 670), (1120, 670), width=4)
    draw_arrow(draw, (1270, 420), (1450, 560), width=4)

    draw.text(
        (90, 1160),
        "代码边界：最大 6 个决策步骤、8 次工具调用、120 秒超时；Research 模式必须通过证据门。",
        fill="#334E68",
        font=font(25),
    )
    image.save(ARCHITECTURE, quality=95)


def set_run_font(run, name="Microsoft YaHei", size=11, bold=None, color=TEXT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman" if name != "Consolas" else name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman" if name != "Consolas" else name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = color


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), GRID)


def style_table(table, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            shade_cell(cell, NAVY if r_idx == 0 else (PALE if r_idx % 2 == 0 else "FFFFFF"))
            if widths and c_idx < len(widths):
                cell.width = Inches(widths[c_idx])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.15
                if c_idx == 0 and len(row.cells) <= 5:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run_font(run, size=9.5, bold=(r_idx == 0), color=RGBColor(255, 255, 255) if r_idx == 0 else TEXT)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = str(header)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    style_table(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Pt(22)

    title = doc.styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = TEXT
    title_ppr = title._element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = TEXT
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("第 ")
    set_run_font(r, size=9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r2 = OxmlElement("w:r")
    t2 = OxmlElement("w:t")
    t2.text = "1"
    r2.append(t2)
    fld.append(r2)
    p._p.append(fld)
    r3 = p.add_run(" 页")
    set_run_font(r3, size=9)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text, *, first_indent=True, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(22) if first_indent else Pt(0)
    if bold_lead and text.startswith(bold_lead):
        run = p.add_run(bold_lead)
        set_run_font(run, bold=True)
        text = text[len(bold_lead):]
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_code(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Pt(12)
    p.paragraph_format.right_indent = Pt(12)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.05
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p._p.get_or_add_pPr().append(shd)
    run = p.add_run(code.rstrip())
    set_run_font(run, name="Consolas", size=8.5)


def append_markdown(doc, path: Path, *, skip_title=True):
    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    skipped = False
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code(doc, "\n".join(code_lines))
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\-:| ]+\|$", lines[i + 1].strip()):
            headers = [x.strip() for x in line.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([x.strip().replace("`", "") for x in lines[i].strip().strip("|").split("|")])
                i += 1
            widths = None
            if len(headers) == 5:
                widths = [0.65, 1.50, 2.25, 1.60, 0.70]
            elif len(headers) == 4:
                widths = [0.6, 2.1, 2.55, 1.55]
            elif len(headers) == 3:
                if headers[0] == "错误码":
                    widths = [1.9, 2.4, 2.7]
                elif headers[0] == "发现的问题":
                    widths = [1.6, 2.6, 2.8]
                elif headers[0] == "变量":
                    widths = [2.2, 2.0, 2.8]
                elif headers[0] == "现象":
                    widths = [1.7, 2.1, 3.2]
                else:
                    widths = [1.6, 2.2, 3.2]
            add_table(doc, headers, rows, widths)
            continue
        if line.startswith("#"):
            hashes = len(line) - len(line.lstrip("#"))
            text = line[hashes:].strip()
            if skip_title and hashes == 1 and not skipped:
                skipped = True
            else:
                if text == "7 主要错误码":
                    doc.add_page_break()
                add_heading(doc, text, min(max(hashes, 1), 3))
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.first_line_indent = Pt(0)
            run = p.add_run(line[2:].replace("`", ""))
            set_run_font(run)
            i += 1
            continue
        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.first_line_indent = Pt(0)
            run = p.add_run(re.sub(r"^\d+\. ", "", line).replace("`", ""))
            set_run_font(run)
            i += 1
            continue
        paragraph_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "- ", "```", "|")) and not re.match(r"^\d+\. ", lines[i]):
            paragraph_lines.append(lines[i].strip())
            i += 1
        text = " ".join(paragraph_lines).replace("**", "").replace("`", "")
        add_body(doc, text)


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    ppr = p._p.get_or_add_pPr()
    for border in ppr.findall(qn("w:pBdr")):
        ppr.remove(border)
    run = p.add_run("AI 信息小屋 Agent 实训报告")
    set_run_font(run, size=24, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.first_line_indent = Pt(0)
    run = p2.add_run("对话式 AI 资讯研究 Agent 的设计 实现与评测")
    set_run_font(run, size=15, bold=False, color=RGBColor(65, 65, 65))
    doc.add_paragraph()
    table = doc.add_table(rows=5, cols=2)
    values = [
        ("项目名称", "AI 信息小屋"),
        ("开发者", "张锦红"),
        ("学号", "________________"),
        ("完成日期", "2026 年 10 月 29 日"),
        ("核心技术", "FastAPI LangGraph Pydantic SQLite React"),
    ]
    for row, (key, value) in zip(table.rows, values):
        row.cells[0].text = key
        row.cells[1].text = value
    style_table(table, [1.6, 4.6])
    doc.add_paragraph()
    p3 = doc.add_paragraph()
    p3.paragraph_format.first_line_indent = Pt(0)
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p3.add_run("本报告依据当前可运行工程、接口实现和自动化测试结果编写。")
    set_run_font(r, size=10, color=RGBColor(80, 80, 80))
    doc.add_page_break()


def add_toc(doc):
    add_heading(doc, "内容目录", 1)
    entries = [
        "一 项目基本信息",
        "二 需求与场景定义",
        "三 Agent 架构设计",
        "四 工具设计",
        "五 Prompt 工程",
        "六 核心实现难点与解法",
        "七 评测与分析",
        "八 总结与反思",
        "附录一 核心源码说明",
        "附录二 部署说明",
    ]
    for entry in entries:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.left_indent = Pt(20)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(entry)
        set_run_font(r, size=11.5)
    doc.add_page_break()


def build_report():
    build_architecture()
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_toc(doc)

    add_heading(doc, "一 项目基本信息", 1)
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["项目名称", "AI 信息小屋"],
            ["开发者", "张锦红"],
            ["完成日期", "2026 年 10 月 29 日"],
            ["运行环境", "Python 3.12 及以上；当前测试环境 Python 3.13.0"],
            ["后端", "FastAPI 0.141.1、LangGraph 1.x、Pydantic 2.13.4、SQLAlchemy 2.0.51"],
            ["前端", "React、TypeScript、Vite"],
            ["模型接口", "OpenAI 兼容 API；当前项目默认配置 DeepSeek"],
            ["代码规模", "约 68,441 行源码与测试文本；421 项自动化测试通过"],
        ],
        [1.4, 5.3],
    )
    add_body(doc, "AI 信息小屋是一套本地运行的 AI 资讯采集、筛选、研究和内容创作系统。本次实训聚焦其中的“灵犀探讯”对话式 Agent。它接收自然语言资讯需求，调用本地资讯库或 AIHOT 发现工具，读取结构化观察结果，再决定继续检索、补充来源、构建证据，或形成带来源限制的回答。")

    add_heading(doc, "二 需求与场景定义", 1)
    add_heading(doc, "2.1 目标用户与使用场景", 2)
    add_body(doc, "目标用户是需要持续跟踪 AI 公司、模型、产品、论文和行业动态的内容创作者。用户通常只会输入“最近 OpenAI 有什么更新”或“只看产品更新”这类不完整需求。Agent 需要先搜索，再允许用户在后续消息中调整时间范围和关注角度，不能因为条件不完整就停止工作。")
    add_heading(doc, "2.2 核心问题", 2)
    bullets = [
        "把自然语言需求转成受限的搜索条件，并在多轮对话中继承已确认的主题。",
        "把本地资讯、AIHOT 外部发现和研究证据区分开，避免把聚合摘要误当成独立核验。",
        "让工具调用可验证、可停止、可审计，阻止未知工具、非法参数和未授权刷新。",
        "在证据不足或事实冲突时明确返回缺口，而不是生成看似完整但无法支持的结论。",
    ]
    for item in bullets:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.first_line_indent = Pt(0)
        set_run_font(p.add_run(item))
    add_heading(doc, "2.3 成功标准", 2)
    add_table(
        doc,
        ["标准", "可观察结果", "验证方式"],
        [
            ["直接搜索", "宽泛需求默认 7 天并立即调用工具", "单元测试 T01 T03"],
            ["多轮继承", "后续修正保留上一轮主题", "接口测试 T04"],
            ["证据门", "单源或冲突内容不能完成 Research", "单元测试 T11 T12"],
            ["安全边界", "非法参数、未知工具和未授权刷新被阻止", "单元测试 T05 T06 T13"],
            ["稳定停止", "步骤、工具数和时间均有上限", "单元测试 T09 T10 与超时用例"],
            ["可运行", "主页、Agent API 和全量测试可通过", "HTTP 健康检查与 421 项测试"],
        ],
        [1.25, 3.4, 2.05],
    )

    add_heading(doc, "三 Agent 架构设计", 1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.add_run().add_picture(str(ARCHITECTURE), width=Inches(6.7))
    cap = doc.add_paragraph("图 1 AI 信息小屋 Agent 的循环结构 工具层与记忆边界")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Pt(0)
    for run in cap.runs:
        set_run_font(run, size=9)
    add_heading(doc, "3.1 感知输入", 2)
    add_body(doc, "Agent 的感知输入包括当前用户消息、同一会话中此前的用户和助手消息、每次工具执行后的 Observation，以及运行模式 query 或 research。API 层只保留角色为 user 或 assistant 且内容非空的历史，避免把内部工具消息跨轮复用。")
    add_heading(doc, "3.2 决策与循环", 2)
    add_body(doc, "LangGraph 图由 decide 和 execute 两个节点组成。decide 调用支持原生 function calling 的模型；模型选择工具时进入 execute，工具结果以 role=tool 的结构化 JSON 写回消息，再回到 decide。模型不再选择工具并且已完成真实资讯搜索时，图才能产生 Answer。")
    add_code(doc, "START -> decide\ndecide --有工具调用--> execute -> decide\ndecide --无工具调用且已搜索--> END\ndecide --达到步骤或超时限制--> END")
    add_heading(doc, "3.3 记忆读写", 2)
    add_body(doc, "会话记忆由 API 层的 OrderedDict 保存，每个会话记录 messages、run_ids、mode 和更新时间；每次新消息都会把历史传入图。工具层另有本轮 seen 映射，只有真正搜索到的文章 ID 才能进入证据工具。两个 Agent 运行不会共享 seen，从而避免跨任务污染证据链。")
    add_body(doc, "当前会话记忆属于进程内短期记忆，最多保留 100 个会话和 100 次运行。SQLite 已用于文章、任务和创作数据，但尚未承担 Agent 会话持久化，因此服务重启后会话会清空。")
    add_heading(doc, "3.4 运行边界", 2)
    add_table(
        doc,
        ["约束", "默认值", "作用"],
        [
            ["最大决策步骤", "6", "阻止模型无限规划"],
            ["最大工具调用", "8", "限制成本和副作用"],
            ["单次超时", "120 秒", "阻止长时间占用同步锁"],
            ["最大输出", "1200 token", "控制回答长度和费用"],
            ["并发方式", "单运行锁", "保证进程内状态一致"],
            ["默认权限", "只读", "刷新信源必须显式授权"],
        ],
        [2.0, 1.4, 3.3],
    )

    add_heading(doc, "四 工具设计", 1)
    append_markdown(doc, ROOT / "02-工具接口说明.md")

    add_heading(doc, "五 Prompt 工程", 1)
    append_markdown(doc, ROOT / "04-Prompt设计文档.md")

    add_heading(doc, "六 核心实现难点与解法", 1)
    add_heading(doc, "6.1 模型可能跳过搜索", 2)
    add_body(doc, "第一个难点是 OpenAI 兼容模型对工具约束的实现不完全一致。即使 System Prompt 要求先搜索，模型仍可能直接回答或先向用户提问。项目没有把这类文本直接返回，而是用 require_search 检查本轮是否存在成功的资讯工具 Observation。没有查询时，图节点追加一条运行时系统消息并继续循环。这样“必须先搜索”成为状态机可验证的条件，而不是只靠提示词自觉遵守。")
    add_heading(doc, "6.2 工具参数和权限可能被模型误用", 2)
    add_body(doc, "第二个难点是模型可能生成错误 JSON、越界 limit、未知字段、未注册工具或臆造文章 ID。项目用 Pydantic strict 模式校验工具参数，未知工具直接返回 UNKNOWN_TOOL，并通过 seen 映射限制证据 ID。刷新信源还要经过 API 的 refresh_sources 布尔值授权；模型文字不能赋予自己权限。")
    add_heading(doc, "6.3 长上下文和跨轮状态容易丢失", 2)
    add_body(doc, "第三个难点是用户常用“更近”“只看产品”“换成 30 天”修正需求。如果只传当前消息，模型会丢失主题。API 层保留用户与助手历史并在下一轮完整传入，同时 System Prompt 要求只更新用户明确改变的条件。工具观察不跨轮直接复用，避免把旧搜索结果当成新一轮证据。")
    add_heading(doc, "6.4 证据数量不等于证据质量", 2)
    add_body(doc, "同一篇稿件被多个转载站引用，并不代表存在多个独立来源。证据构建按 source_identity 去重，同时检查一手来源、待核实内容和数值冲突。Research 模式只有 evidence_ready 为真才能完成；否则返回 needs_evidence 和明确缺口。")
    add_heading(doc, "6.5 外部扩展必须可移除", 2)
    add_body(doc, "AIHOT 以 AihotResearchExtension 独立模块接入。核心 ResearchTools 只接收扩展定义，不依赖 AIHOT 具体实现。关闭 WECHAT_AGENT_AIHOT_ENABLED 或删除适配器注册即可移除，不影响本地资讯搜索和证据工具。")

    add_heading(doc, "七 评测与分析", 1)
    append_markdown(doc, ROOT / "05-Agent评测记录.md")

    add_heading(doc, "八 总结与反思", 1)
    add_heading(doc, "8.1 技术收获", 2)
    add_body(doc, "本项目把 Agent 从“模型回复接口”拆成可观测的状态机。模型负责选择下一步，代码负责参数校验、权限、停止条件、证据门和状态转换。这样可以用脚本模型固定决策轨迹，验证每一条分支，而不依赖付费模型的随机输出。")
    add_heading(doc, "8.2 Agent 与普通 API 调用的区别", 2)
    add_body(doc, "普通 API 调用通常是一次输入对应一次输出，流程由调用方预先写死。此 Agent 会在运行中读取 Observation，再决定下一项工具或结束回答，工具次数和路径由目标与结果共同决定。它仍然不是无限自治系统，因为所有工具、参数、权限、步骤和完成条件都由服务端控制。")
    add_heading(doc, "8.3 后续优化方向", 2)
    for item in [
        "把 Agent 会话和运行记录持久化到 SQLite，使服务重启后可以恢复。",
        "把同步单运行锁升级为队列任务，增加运行状态流式推送和取消能力。",
        "增加定时真实链路评测，监测 AIHOT、RSS 和模型服务商的可用性与响应格式变化。",
        "建立基于人工标注问题集的回答质量评测，覆盖来源引用正确性、事件聚合和时间范围解释。",
        "在证据构建阶段加入原文级抓取与段落引用，减少只依赖摘要匹配的限制。",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.first_line_indent = Pt(0)
        set_run_font(p.add_run(item))

    appendix_one = add_heading(doc, "附录一 核心源码说明", 1)
    appendix_one.paragraph_format.page_break_before = True
    append_markdown(doc, ROOT / "03-核心源码说明.md")
    appendix_two = add_heading(doc, "附录二 部署说明", 1)
    appendix_two.paragraph_format.page_break_before = True
    append_markdown(doc, ROOT / "06-部署说明.md")

    doc.core_properties.title = "AI 信息小屋 Agent 实训报告"
    doc.core_properties.subject = "Agent 架构 工具 Prompt 评测与部署"
    doc.core_properties.author = "张锦红"
    doc.core_properties.keywords = "Agent, LangGraph, FastAPI, AI资讯, 实训报告"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()

